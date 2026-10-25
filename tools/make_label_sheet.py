"""Build the human-labelling workbook (.xlsx) with dropdowns, ready to upload to Google Sheets.

Sheet "labels": every article, assigned round-robin to one labeller.
Sheet "overlap": 30 of those articles again, for a second labeller who must not look at
"labels". The two sheets are compared later to report inter-annotator agreement.

Dropdowns:
- label_primary_ticker: only that row's candidate TICKERS (leave blank for no single subject)
- label_event_type: the nine event types used by the system
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

EVENT_TYPES = ["earnings", "guidance", "m&a", "rating", "contract", "product",
               "management", "regulatory", "other"]
PEOPLE = ["Mingyang", "Wenhao", "Yuhua", "Zhaoyang"]
HEADER = ["id", "title", "content", "tickers", "url", "labeller",
          "label_primary_ticker", "label_event_type", "notes"]
WIDTHS = [22, 50, 80, 18, 30, 12, 20, 18, 30]
PRIMARY_COL, EVENT_COL = "G", "H"


def add_sheet(wb: Workbook, title: str, rows: list[tuple[dict, str]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "B2"
    for letter, width in zip("ABCDEFGHI", WIDTHS):
        ws.column_dimensions[letter].width = width

    events = DataValidation(type="list", formula1='"' + ",".join(EVENT_TYPES) + '"',
                            allow_blank=False, showErrorMessage=True,
                            errorTitle="Invalid event type", error="Choose one of the nine event types.")
    ws.add_data_validation(events)

    for row_number, (article, labeller) in enumerate(rows, start=2):
        ws.append([article["ID"], article["TITLE"], article["CONTENT"], " ".join(article["TICKERS"]),
                   article["URL"], labeller, "", "", ""])
        for cell in ws[row_number]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        events.add(f"{EVENT_COL}{row_number}")
        # Per-row ticker dropdown: only this article's candidates. Blank is allowed and means
        # "no single primary company".
        tickers = DataValidation(type="list", formula1='"' + ",".join(article["TICKERS"]) + '"',
                                 allow_blank=True, showErrorMessage=True,
                                 errorTitle="Not a candidate ticker",
                                 error="Pick a ticker from this row's list, or leave blank.")
        ws.add_data_validation(tickers)
        tickers.add(f"{PRIMARY_COL}{row_number}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="data/news/news_eval.jsonl")
    parser.add_argument("--output", default="data/news/news_eval_labelling.xlsx")
    parser.add_argument("--overlap", type=int, default=30)
    parser.add_argument("--seed", type=int, default=6725)
    args = parser.parse_args(argv)

    articles = [json.loads(line) for line in Path(args.input).read_text(encoding="utf-8").splitlines() if line.strip()]
    for article in articles:
        if len(",".join(article["TICKERS"])) > 250:
            raise ValueError(f"{article['ID']}: too many tickers for a dropdown")

    first = [(article, PEOPLE[i % len(PEOPLE)]) for i, article in enumerate(articles)]
    rng = random.Random(args.seed)
    picked = sorted(rng.sample(range(len(articles)), min(args.overlap, len(articles))))
    # The second labeller is always a different person from the first.
    second = [(articles[i], PEOPLE[(i + 1) % len(PEOPLE)]) for i in picked]

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "labels", first)
    add_sheet(wb, "overlap", second)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(json.dumps({"articles": len(first), "overlap": len(second), "output": args.output}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
